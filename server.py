# -*- coding: utf-8 -*-
"""PEA transformer field-collection server."""
from __future__ import annotations

import json
import os
import re
import secrets
import shutil
import socket
import subprocess
import sys
import tempfile
import threading
import time
import uuid
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from io import BytesIO

from flask import Flask, jsonify, request, send_from_directory, send_file, abort, redirect, url_for
from werkzeug.middleware.proxy_fix import ProxyFix

ROOT = Path(__file__).resolve().parent
DATA_DIR = ROOT / "data"
UPLOAD_DIR = DATA_DIR / "uploads"
LOCATIONS_FILE = DATA_DIR / "locations.json"
SUBMISSIONS_FILE = DATA_DIR / "submissions.json"
CONFIG_FILE = DATA_DIR / "config.json"
EXCEL_FILE = ROOT / "Rak-D.xlsx"
HTML_FILE = ROOT / "trpjk.html"

HOST = os.environ.get("PEA_HOST", "0.0.0.0")
PORT = int(os.environ.get("PORT") or os.environ.get("PEA_PORT", "5050"))
ALLOWED_IMAGE_EXT = {".jpg", ".jpeg", ".png", ".webp", ".gif", ".heic", ".heif"}
MAX_CONTENT = 32 * 1024 * 1024

BLOCKED_TUNNEL_HOSTS = {
    "admin.localhost.run",
    "www.localhost.run",
    "docs.localhost.run",
    "localhost.run",
}
ANSI_RE = re.compile(r"\x1b(?:\[[0-9;]*[A-Za-z]|\]8;;[^\x1b]*\x1b\\)")
TUNNEL_HOST_RE = re.compile(
    r"(?:https://)?([a-zA-Z0-9](?:[a-zA-Z0-9-]{0,61}[a-zA-Z0-9])?\.(?:lhr\.life|lhr\.rocks|a\.pinggy\.link|pinggy\.link))",
    re.I,
)


def is_real_tunnel_base(url: str) -> bool:
    if not url:
        return False
    host = url.replace("https://", "").replace("http://", "").split("/")[0].strip().lower()
    if ":" in host:
        host = host.split(":", 1)[0]
    if not host or host in BLOCKED_TUNNEL_HOSTS:
        return False
    if "localhost.run" in host:
        return False
    return (
        host.endswith(".lhr.life")
        or host.endswith(".lhr.rocks")
        or host.endswith(".pinggy.link")
    )


def extract_tunnel_bases(text: str) -> list[str]:
    cleaned = ANSI_RE.sub(" ", text or "")
    found: list[str] = []
    for host in TUNNEL_HOST_RE.findall(cleaned):
        base = "https://" + host.lower().rstrip(".")
        if is_real_tunnel_base(base) and base not in found:
            found.append(base)
    return found


lock = threading.Lock()
PUBLIC_BASE = {"url": None, "ok": False}
TUNNEL_RESTART = threading.Event()


def tunnel_key_dir() -> Path:
    base = os.environ.get("LOCALAPPDATA") or os.environ.get("TEMP") or str(DATA_DIR)
    folder = Path(base) / "pea-transformer"
    folder.mkdir(parents=True, exist_ok=True)
    return folder


TUNNEL_KEY = None  # set on first use
TUNNEL_PUB = None

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = MAX_CONTENT
app.wsgi_app = ProxyFix(app.wsgi_app, x_for=1, x_proto=1, x_host=1)


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()


def ensure_dirs() -> None:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    UPLOAD_DIR.mkdir(parents=True, exist_ok=True)


def read_json(path: Path, default):
    if not path.exists():
        return default
    with path.open(encoding="utf-8") as f:
        return json.load(f)


def write_json(path: Path, payload) -> None:
    tmp = path.with_suffix(".tmp")
    with tmp.open("w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=2)
    tmp.replace(path)


def load_excel_locations() -> list[dict]:
    import openpyxl

    wb = openpyxl.load_workbook(EXCEL_FILE, data_only=True)
    ws = wb["WorkOrders"]
    rows = []
    for row in ws.iter_rows(min_row=3, values_only=True):
        rid, wo, addr, coords, meter, tr = row[0], row[1], row[2], row[3], row[4], row[5]
        if rid is None or not coords:
            continue
        lat_s, lng_s = str(coords).split(",")
        rows.append(
            {
                "id": int(rid),
                "wo": wo or "",
                "address": addr or "",
                "lat": float(lat_s.strip()),
                "lng": float(lng_s.strip()),
                "meter": meter or "",
                "transformer": tr or "",
                "assignee": row[14] or "",
                "queue": row[15] if row[15] is not None else "",
            }
        )
    return rows


def sync_locations() -> list[dict]:
    ensure_dirs()
    if EXCEL_FILE.exists():
        try:
            rows = load_excel_locations()
            write_json(LOCATIONS_FILE, rows)
            return rows
        except Exception as exc:
            print("Excel sync failed, using cache:", exc)
    cached = read_json(LOCATIONS_FILE, [])
    if not cached:
        raise RuntimeError("ไม่พบข้อมูลหม้อแปลง (Rak-D.xlsx หรือ data/locations.json)")
    return cached


def load_config() -> dict:
    ensure_dirs()
    cfg = read_json(CONFIG_FILE, {})
    if not cfg.get("token"):
        cfg["token"] = secrets.token_urlsafe(9).replace("_", "").replace("-", "")[:12]
        cfg["createdAt"] = utc_now()
        write_json(CONFIG_FILE, cfg)
    if cfg.get("stableBaseUrl") and not is_real_tunnel_base(cfg.get("stableBaseUrl") or ""):
        cfg.pop("stableBaseUrl", None)
        write_json(CONFIG_FILE, cfg)
    return cfg


def is_hosted() -> bool:
    return bool(
        os.environ.get("RENDER")
        or os.environ.get("RAILWAY_ENVIRONMENT")
        or os.environ.get("FLY_APP_NAME")
        or os.environ.get("PUBLIC_URL")
    )


def hosted_public_url() -> str | None:
    cfg = load_config()
    path = f"/f/{cfg['token']}"
    env_url = (os.environ.get("PUBLIC_URL") or "").rstrip("/")
    if env_url:
        return env_url + path
    try:
        origin = request.host_url.rstrip("/")
        proto = request.headers.get("X-Forwarded-Proto")
        if proto:
            origin = proto.split(",")[0].strip() + "://" + request.host
        return origin + path
    except RuntimeError:
        return None


def lan_ip() -> str:
    try:
        sock = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
        sock.connect(("8.8.8.8", 80))
        ip = sock.getsockname()[0]
        sock.close()
        return ip
    except OSError:
        return "127.0.0.1"


def feeder_complete(feeder: dict) -> bool:
    def filled(key: str) -> bool:
        val = feeder.get(key)
        return val is not None and str(val).strip() != ""

    return filled("ia") and filled("ib") and filled("ic")


def location_status(entry: dict | None) -> str:
    if not entry:
        return "pending"
    if entry.get("finished"):
        return "completed"
    images = entry.get("images") or []
    feeders = entry.get("feeders") or []
    notes = (entry.get("notes") or "").strip()
    has_current = any(feeder_complete(f) for f in feeders)
    has_partial_current = any(
        str(f.get(k) or "").strip() for f in feeders for k in ("ia", "ib", "ic")
    )
    if images or notes or has_partial_current or has_current:
        return "partial"
    return "pending"


def submissions() -> dict:
    return read_json(SUBMISSIONS_FILE, {})


def save_submissions(data: dict) -> None:
    write_json(SUBMISSIONS_FILE, data)


def loc_key(loc_id) -> str:
    return str(loc_id)


STATUS_LABEL = {
    "pending": "รอดำเนินการ",
    "partial": "ลงข้อมูลไม่ครบ",
    "completed": "ดำเนินการแล้ว",
}

THUMB_MAX_W = 280
THUMB_MAX_H = 210


def _ensure_pillow() -> bool:
    try:
        from PIL import Image as PILImage
    except ImportError:
        try:
            subprocess.run(
                [sys.executable, "-m", "pip", "install", "Pillow>=10.0.0"],
                check=True,
                capture_output=True,
                text=True,
                creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
            )
            from PIL import Image as PILImage
        except Exception as exc:
            print("Pillow is required to embed photos in Excel:", exc)
            return False
    try:
        import openpyxl.drawing.image as xl_image

        if not getattr(xl_image, "PILImage", None):
            xl_image.PILImage = PILImage
    except Exception:
        pass
    return True


def _safe_filename(text: str, fallback: str = "file") -> str:
    cleaned = re.sub(r'[\\/:*?"<>|]+', "_", str(text or "").strip())
    cleaned = re.sub(r"\s+", " ", cleaned).strip(" ._")
    return (cleaned or fallback)[:80]


def _photo_folder(loc: dict) -> str:
    assignee = _safe_filename(loc.get("assignee") or "", "ยังไม่ระบุผู้รับผิดชอบ")
    tr = _safe_filename(loc.get("transformer") or "", f"ID{loc.get('id')}")
    addr = _safe_filename(loc.get("address") or "", "")
    point = f"{tr}_{addr}" if addr else tr
    return f"รูป/{assignee}/{point}"


def _write_excel_thumb(src: Path, dest: Path) -> bool:
    if not src.is_file():
        return False
    dest.parent.mkdir(parents=True, exist_ok=True)
    try:
        from PIL import Image as PILImage

        with PILImage.open(src) as im:
            im = im.convert("RGB")
            im.thumbnail((THUMB_MAX_W, THUMB_MAX_H))
            im.save(dest, format="JPEG", quality=88)
        return dest.is_file() and dest.stat().st_size > 0
    except Exception:
        if src.suffix.lower() in {".jpg", ".jpeg", ".png", ".gif"}:
            try:
                shutil.copyfile(src, dest)
                return True
            except OSError:
                return False
        return False


def _original_export_name(img_rec: dict, src: Path, idx: int, loc: dict | None = None) -> str:
    original = _safe_filename(img_rec.get("original") or src.name, src.name)
    if not Path(original).suffix:
        original += src.suffix or ".jpg"
    tr = _safe_filename((loc or {}).get("transformer") or "", "TR")
    return f"{tr}_{idx:02d}_{original}"


def build_export_xlsx(tmp_dir: Path) -> tuple[BytesIO, list[tuple[str, Path]]]:
    _ensure_pillow()
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    with lock:
        locs = read_json(LOCATIONS_FILE, [])
        saved = submissions()

    wb = Workbook()
    ws = wb.active
    ws.title = "ข้อมูลหม้อแปลง"

    headers = [
        "ลำดับ",
        "ID",
        "หม้อแปลง",
        "Work Order",
        "ผู้รับผิดชอบ",
        "ที่อยู่",
        "Lat",
        "Lng",
        "สถานะ",
        "หมายเหตุ",
        "ฟีดเดอร์ / กระแส (A)",
        "ผู้บันทึก",
        "อัปเดตล่าสุด",
        "จำนวนรูป",
        "โฟลเดอร์รูปต้นฉบับ",
    ]
    max_images = 0
    rows_data = []
    originals: list[tuple[str, Path]] = []
    for loc in locs:
        entry = saved.get(loc_key(loc["id"]), {})
        images = entry.get("images") or []
        max_images = max(max_images, len(images))
        rows_data.append((loc, entry, images))

    for i in range(1, max_images + 1):
        headers.append(f"รูป {i}")

    header_font = Font(bold=True)
    for col, title in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.font = header_font
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    ws.freeze_panes = "A2"
    img_col_start = len(headers) - max_images + 1 if max_images else len(headers) + 1

    for idx, (loc, entry, images) in enumerate(rows_data, 1):
        row = idx + 1
        status = location_status(entry)
        feeders = entry.get("feeders") or []
        feeder_lines = []
        for f in feeders:
            name = str(f.get("name") or "").strip() or "ฟีดเดอร์"
            feeder_lines.append(
                f"{name}: Ia={f.get('ia', '') or '-'} Ib={f.get('ib', '') or '-'} Ic={f.get('ic', '') or '-'}"
            )
        feeder_text = "\n".join(feeder_lines)
        folder = _photo_folder(loc) if images else ""

        values = [
            idx,
            loc.get("id"),
            loc.get("transformer") or "",
            loc.get("wo") or "",
            loc.get("assignee") or "",
            loc.get("address") or "",
            loc.get("lat"),
            loc.get("lng"),
            STATUS_LABEL.get(status, status),
            entry.get("notes") or "",
            feeder_text,
            entry.get("updatedBy") or "",
            entry.get("updatedAt") or entry.get("finishedAt") or "",
            len(images),
            folder,
        ]
        for col, value in enumerate(values, 1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)

        row_height = 90
        for img_idx, img_rec in enumerate(images):
            filename = img_rec.get("filename") or ""
            col_letter = get_column_letter(img_col_start + img_idx)
            if not filename:
                continue
            src_path = UPLOAD_DIR / str(loc["id"]) / filename
            export_name = _original_export_name(img_rec, src_path, img_idx + 1, loc)
            if not src_path.is_file():
                ws.cell(row=row, column=img_col_start + img_idx, value=f"(ไม่พบไฟล์: {filename})")
                continue
            originals.append((f"{folder}/{export_name}", src_path))
            ws.cell(row=row, column=img_col_start + img_idx, value=export_name)
            thumb_path = tmp_dir / f"r{loc['id']}_{img_idx + 1}.jpg"
            if not _write_excel_thumb(src_path, thumb_path):
                ws.cell(row=row, column=img_col_start + img_idx, value=f"(เปิดรูปไม่ได้: {export_name})")
                continue
            try:
                xl_img = XLImage(str(thumb_path))
                if xl_img.width and xl_img.height:
                    scale = min(THUMB_MAX_W / xl_img.width, THUMB_MAX_H / xl_img.height, 1.0)
                    xl_img.width = max(1, int(xl_img.width * scale))
                    xl_img.height = max(1, int(xl_img.height * scale))
                else:
                    xl_img.width = THUMB_MAX_W
                    xl_img.height = THUMB_MAX_H
                ws.add_image(xl_img, f"{col_letter}{row}")
                row_height = max(row_height, xl_img.height * 0.75 + 16)
                ws.column_dimensions[col_letter].width = 22
            except Exception:
                ws.cell(row=row, column=img_col_start + img_idx, value=f"(แนบรูปไม่ได้: {export_name})")

        if images:
            ws.row_dimensions[row].height = row_height

    widths = [6, 6, 18, 12, 14, 28, 10, 10, 14, 24, 28, 12, 20, 8, 28]
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width

    ws2 = wb.create_sheet("ฟีดเดอร์")
    feeder_headers = [
        "ID",
        "หม้อแปลง",
        "ผู้รับผิดชอบ",
        "ชื่อฟีดเดอร์",
        "Ia (A)",
        "Ib (A)",
        "Ic (A)",
    ]
    for col, title in enumerate(feeder_headers, 1):
        cell = ws2.cell(row=1, column=col, value=title)
        cell.font = header_font
    frow = 2
    for loc, entry, _ in rows_data:
        feeders = entry.get("feeders") or []
        if not feeders:
            continue
        for f in feeders:
            ws2.cell(row=frow, column=1, value=loc.get("id"))
            ws2.cell(row=frow, column=2, value=loc.get("transformer") or "")
            ws2.cell(row=frow, column=3, value=loc.get("assignee") or "")
            ws2.cell(row=frow, column=4, value=f.get("name") or "")
            ws2.cell(row=frow, column=5, value=f.get("ia") or "")
            ws2.cell(row=frow, column=6, value=f.get("ib") or "")
            ws2.cell(row=frow, column=7, value=f.get("ic") or "")
            frow += 1

    ws3 = wb.create_sheet("รายการรูป")
    photo_headers = ["ID", "หม้อแปลง", "ผู้รับผิดชอบ", "ชื่อไฟล์ใน ZIP", "เส้นทางใน ZIP"]
    for col, title in enumerate(photo_headers, 1):
        cell = ws3.cell(row=1, column=col, value=title)
        cell.font = header_font
    prow = 2
    for zip_rel, src_path in originals:
        loc_id = src_path.parent.name
        loc = next((x for x in locs if str(x.get("id")) == loc_id), {})
        ws3.cell(row=prow, column=1, value=loc.get("id") or loc_id)
        ws3.cell(row=prow, column=2, value=loc.get("transformer") or "")
        ws3.cell(row=prow, column=3, value=loc.get("assignee") or "")
        ws3.cell(row=prow, column=4, value=Path(zip_rel).name)
        ws3.cell(row=prow, column=5, value=zip_rel)
        prow += 1
    for col, width in enumerate([8, 18, 16, 36, 48], 1):
        ws3.column_dimensions[get_column_letter(col)].width = width

    bio = BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio, originals


def build_export_zip() -> tuple[BytesIO, str]:
    stamp = datetime.now().strftime("%Y%m%d_%H%M")
    xlsx_name = f"PEA_ข้อมูลหม้อแปลง_{stamp}.xlsx"
    zip_name = f"PEA_ข้อมูลหม้อแปลง_{stamp}.zip"
    tmp_dir = Path(tempfile.mkdtemp(prefix="pea-export-"))
    try:
        xlsx_bio, originals = build_export_xlsx(tmp_dir)
        zbio = BytesIO()
        with zipfile.ZipFile(zbio, "w", compression=zipfile.ZIP_DEFLATED) as zf:
            zf.writestr(xlsx_name, xlsx_bio.getvalue())
            used = set()
            for zip_rel, src_path in originals:
                name = zip_rel
                if name in used:
                    stem = Path(zip_rel).stem
                    suffix = Path(zip_rel).suffix
                    n = 2
                    while f"{Path(zip_rel).parent.as_posix()}/{stem}_{n}{suffix}" in used:
                        n += 1
                    name = f"{Path(zip_rel).parent.as_posix()}/{stem}_{n}{suffix}"
                used.add(name)
                zf.write(src_path, name)
        zbio.seek(0)
        return zbio, zip_name
    finally:
        shutil.rmtree(tmp_dir, ignore_errors=True)


def require_token() -> None:
    cfg = load_config()
    token = request.view_args.get("token") if request.view_args else None
    if not token:
        token = request.args.get("k") or request.headers.get("X-Share-Token")
    if token != cfg["token"]:
        abort(403, description="ลิงก์ไม่ถูกต้อง")


@app.route("/")
def root():
    cfg = load_config()
    return redirect(url_for("app_page", token=cfg["token"]))


@app.route("/f/<token>")
def app_page(token):
    cfg = load_config()
    if token != cfg["token"]:
        abort(403)
    return send_from_directory(ROOT, HTML_FILE.name)


def public_share_url() -> str | None:
    base = PUBLIC_BASE.get("url")
    if not base or not is_real_tunnel_base(base):
        return None
    cfg = load_config()
    return f"{base.rstrip('/')}/f/{cfg['token']}"


TUNNEL_ADMIN = "https://admin.localhost.run/"


def _tunnel_paths() -> tuple[Path, Path]:
    global TUNNEL_KEY, TUNNEL_PUB
    if TUNNEL_KEY is None:
        d = tunnel_key_dir()
        TUNNEL_KEY = d / "tunnel_key"
        TUNNEL_PUB = d / "tunnel_key.pub"
    return TUNNEL_KEY, TUNNEL_PUB


def ensure_tunnel_key() -> None:
    key, pub = _tunnel_paths()
    if key.exists():
        return
    old_key = DATA_DIR / "tunnel_key"
    old_pub = DATA_DIR / "tunnel_key.pub"
    if old_key.exists():
        shutil.copy2(old_key, key)
        if old_pub.exists():
            shutil.copy2(old_pub, pub)
        return
    subprocess.run(
        ["ssh-keygen", "-t", "ed25519", "-f", str(key), "-N", "", "-q"],
        check=True,
        creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
    )
    try:
        os.chmod(key, 0o600)
    except OSError:
        pass


def tunnel_public_key() -> str:
    ensure_tunnel_key()
    _, pub = _tunnel_paths()
    return pub.read_text(encoding="utf-8").strip()


def mark_stable_tunnel(base_url: str) -> None:
    if not is_real_tunnel_base(base_url):
        return
    cfg = load_config()
    base = base_url.rstrip("/")
    if cfg.get("stableBaseUrl") == base:
        return
    cfg["stableBaseUrl"] = base
    cfg["tunnelKeyRegistered"] = True
    cfg["stableSince"] = utc_now()
    write_json(CONFIG_FILE, cfg)


def effective_public_url() -> str | None:
    live = public_share_url()
    if live:
        return live
    return None


def make_share_url(base: str) -> str:
    cfg = load_config()
    return f"{base.rstrip('/')}/f/{cfg['token']}"


def probe_public(url: str) -> bool:
    if not url or not url.startswith("http"):
        return False
    if not is_real_tunnel_base(url):
        return False
    try:
        import urllib.request

        req = urllib.request.Request(
            url,
            headers={"Cache-Control": "no-cache", "Pragma": "no-cache", "User-Agent": "PEA-app"},
        )
        with urllib.request.urlopen(req, timeout=10) as resp:
            status = getattr(resp, "status", 200)
            body = resp.read(12000).decode("utf-8", "replace").lower()
            if status >= 400:
                return False
            if "no tunnel here" in body or "tunnel not found" in body:
                return False
            if "your billing needs attention" in body:
                return False
            return True
    except Exception:
        return False


def probe_with_retry(url: str, tries: int = 4, delay: float = 1.2) -> bool:
    for _ in range(tries):
        if probe_public(url):
            return True
        time.sleep(delay)
    return False


def clear_public_link() -> None:
    PUBLIC_BASE["url"] = None
    PUBLIC_BASE["ok"] = False
    save_public_url(None)


def load_saved_public_url() -> str | None:
    path = DATA_DIR / "public_url.txt"
    if not path.exists():
        return None
    text = path.read_text(encoding="utf-8").strip()
    return text or None


def save_public_url(url: str | None) -> None:
    path = DATA_DIR / "public_url.txt"
    if url:
        path.write_text(url, encoding="utf-8")
    elif path.exists():
        path.unlink(missing_ok=True)


@app.before_request
def _honor_forwarded_proto():
    proto = request.headers.get("X-Forwarded-Proto")
    if proto:
        request.environ["wsgi.url_scheme"] = proto.split(",")[0].strip()


@app.after_request
def _no_cache_html_and_meta(resp):
    path = request.path or ""
    if resp.mimetype == "text/html" or path.endswith("/api/meta") or "/api/meta" in path:
        resp.headers["Cache-Control"] = "no-store, no-cache, must-revalidate"
        resp.headers["Pragma"] = "no-cache"
    return resp


@app.route("/api/meta")
@app.route("/f/<token>/api/meta")
def api_meta(token=None):
    if token:
        require_token()
    cfg = load_config()
    ip = lan_ip()
    path = f"/f/{cfg['token']}"
    hosted = is_hosted()
    if hosted:
        display = hosted_public_url()
        return jsonify(
            {
                "token": cfg["token"],
                "localUrl": f"http://127.0.0.1:{PORT}{path}",
                "lanUrl": f"http://{ip}:{PORT}{path}",
                "publicUrl": display,
                "livePublicUrl": display,
                "publicOk": True,
                "stableBaseUrl": display,
                "tunnelReady": True,
                "tunnelRegistered": True,
                "tunnelPublicKey": "",
                "tunnelAdminUrl": "",
                "linkPermanent": True,
                "hosted": True,
                "tunnelNote": "ลิงก์คงที่บนคลาวด์ · รูปที่อัปโหลดอาจหายเมื่อ Render รีสตาร์ท (แผนฟรี) ควรกดส่งออก Excel สำรองเป็นประจำ",
                "port": PORT,
            }
        )
    registered = bool(cfg.get("tunnelKeyRegistered"))
    stable = (cfg.get("stableBaseUrl") or "").rstrip("/")
    live = public_share_url() if PUBLIC_BASE.get("ok") else None
    live_matches = bool(live and stable and live.startswith(stable))
    if live:
        note = "ส่งลิงก์นี้ในไลน์ได้เลย ต้องเปิด start.bat ค้างไว้ ห้ามปิดฝาโน้ตบุ๊ค"
    else:
        note = "กำลังสร้างลิงก์สาธารณะ รอสักครู่ ถ้ามือถือกับคอมอยู่ Wi-Fi เดียวกันใช้ลิงก์ในเครือข่ายได้"
    resp = jsonify(
        {
            "token": cfg["token"],
            "localUrl": f"http://127.0.0.1:{PORT}{path}",
            "lanUrl": f"http://{ip}:{PORT}{path}",
            "publicUrl": live,
            "livePublicUrl": live,
            "publicOk": bool(live),
            "stableBaseUrl": stable or None,
            "tunnelReady": bool(live),
            "tunnelRegistered": registered,
            "tunnelPublicKey": "",
            "tunnelAdminUrl": "",
            "linkPermanent": registered and live_matches,
            "hosted": False,
            "tunnelNote": note,
            "port": PORT,
        }
    )
    resp.headers["Cache-Control"] = "no-store, no-cache, must-revalidate"
    resp.headers["Pragma"] = "no-cache"
    return resp


@app.route("/api/tunnel/confirm", methods=["POST"])
@app.route("/f/<token>/api/tunnel/confirm", methods=["POST"])
def api_tunnel_confirm(token=None):
    if token:
        require_token()
    ensure_tunnel_key()
    cfg = load_config()
    cfg["tunnelKeyRegistered"] = True
    write_json(CONFIG_FILE, cfg)
    TUNNEL_RESTART.set()
    return jsonify({
        "ok": True,
        "restartRequired": True,
        "message": "กำลังต่อด้วยคีย์ SSH — รอสักครู่ ลิงก์อาจเปลี่ยนครั้งนี้ครั้งเดียว จากนั้นจะเป็นอันเดิมทุกครั้งที่เปิด start.bat",
    })


@app.route("/api/locations")
@app.route("/f/<token>/api/locations")
def api_locations(token=None):
    if token:
        require_token()
    with lock:
        locs = read_json(LOCATIONS_FILE, [])
        saved = submissions()
    out = []
    for loc in locs:
        entry = saved.get(loc_key(loc["id"]), {})
        out.append(
            {
                **loc,
                "data": {
                    "notes": entry.get("notes", ""),
                    "feeders": entry.get("feeders", []),
                    "images": entry.get("images", []),
                    "finished": bool(entry.get("finished")),
                    "updatedAt": entry.get("updatedAt"),
                    "updatedBy": entry.get("updatedBy", ""),
                },
                "status": location_status(entry),
            }
        )
    return jsonify(out)


@app.route("/api/locations/<int:loc_id>", methods=["PUT"])
@app.route("/f/<token>/api/locations/<int:loc_id>", methods=["PUT"])
def api_save_location(loc_id, token=None):
    if token:
        require_token()
    payload = request.get_json(silent=True) or {}
    feeders = payload.get("feeders") or []
    cleaned = []
    for feeder in feeders:
        cleaned.append(
            {
                "name": str(feeder.get("name") or "").strip(),
                "ia": "" if feeder.get("ia") in (None, "") else str(feeder.get("ia")).strip(),
                "ib": "" if feeder.get("ib") in (None, "") else str(feeder.get("ib")).strip(),
                "ic": "" if feeder.get("ic") in (None, "") else str(feeder.get("ic")).strip(),
            }
        )
    with lock:
        locs = {int(x["id"]) for x in read_json(LOCATIONS_FILE, [])}
        if loc_id not in locs:
            abort(404)
        saved = submissions()
        key = loc_key(loc_id)
        current = saved.get(key, {})
        current["notes"] = str(payload.get("notes") or "")
        current["feeders"] = cleaned
        current["images"] = current.get("images") or []
        if payload.get("finished") is True:
            current["finished"] = True
            current["finishedAt"] = utc_now()
        elif payload.get("finished") is False:
            current["finished"] = False
            current.pop("finishedAt", None)
        current["updatedAt"] = utc_now()
        current["updatedBy"] = str(payload.get("updatedBy") or "")[:80]
        saved[key] = current
        save_submissions(saved)
        status = location_status(current)
    return jsonify({"ok": True, "status": status, "data": current})


@app.route("/api/locations/<int:loc_id>/reopen", methods=["POST"])
@app.route("/f/<token>/api/locations/<int:loc_id>/reopen", methods=["POST"])
def api_reopen_location(loc_id, token=None):
    if token:
        require_token()
    with lock:
        locs = {int(x["id"]) for x in read_json(LOCATIONS_FILE, [])}
        if loc_id not in locs:
            abort(404)
        saved = submissions()
        key = loc_key(loc_id)
        current = saved.get(key) or {"notes": "", "feeders": [], "images": []}
        current["finished"] = False
        current.pop("finishedAt", None)
        current["updatedAt"] = utc_now()
        saved[key] = current
        save_submissions(saved)
        status = location_status(current)
    return jsonify({"ok": True, "status": status, "data": current})


@app.route("/api/locations/<int:loc_id>/images", methods=["POST"])
@app.route("/f/<token>/api/locations/<int:loc_id>/images", methods=["POST"])
def api_upload_images(loc_id, token=None):
    if token:
        require_token()
    files = request.files.getlist("images")
    if not files:
        abort(400, description="ไม่พบไฟล์รูป")
    folder = UPLOAD_DIR / str(loc_id)
    folder.mkdir(parents=True, exist_ok=True)
    added = []
    with lock:
        locs = {int(x["id"]) for x in read_json(LOCATIONS_FILE, [])}
        if loc_id not in locs:
            abort(404)
        saved = submissions()
        key = loc_key(loc_id)
        current = saved.get(key, {"notes": "", "feeders": [], "images": []})
        current.setdefault("images", [])
        for f in files:
            ext = Path(f.filename or "").suffix.lower() or ".jpg"
            if ext not in ALLOWED_IMAGE_EXT:
                ext = ".jpg"
            name = f"{uuid.uuid4().hex}{ext}"
            dest = folder / name
            f.save(dest)
            rec = {"filename": name, "original": Path(f.filename or name).name}
            current["images"].append(rec)
            added.append(rec)
        current["updatedAt"] = utc_now()
        saved[key] = current
        save_submissions(saved)
        status = location_status(current)
    return jsonify({"ok": True, "images": added, "status": status})


@app.route("/api/locations/<int:loc_id>/images/<filename>", methods=["DELETE"])
@app.route("/f/<token>/api/locations/<int:loc_id>/images/<filename>", methods=["DELETE"])
def api_delete_image(loc_id, filename, token=None):
    if token:
        require_token()
    if "/" in filename or "\\" in filename:
        abort(400)
    with lock:
        saved = submissions()
        key = loc_key(loc_id)
        current = saved.get(key)
        if not current:
            abort(404)
        before = len(current.get("images") or [])
        current["images"] = [img for img in current.get("images") or [] if img.get("filename") != filename]
        if len(current["images"]) == before:
            abort(404)
        current["updatedAt"] = utc_now()
        saved[key] = current
        save_submissions(saved)
        path = UPLOAD_DIR / str(loc_id) / filename
        if path.exists():
            path.unlink()
        status = location_status(current)
    return jsonify({"ok": True, "status": status})


@app.route("/api/export")
@app.route("/f/<token>/api/export")
def api_export(token=None):
    if token:
        require_token()
    bio, filename = build_export_zip()
    return send_file(
        bio,
        mimetype="application/zip",
        as_attachment=True,
        download_name=filename,
    )


@app.route("/uploads/<int:loc_id>/<path:filename>")
def serve_upload(loc_id, filename):
    folder = UPLOAD_DIR / str(loc_id)
    return send_from_directory(folder, filename)


def _open_shared_log(path: Path):
    """Write handle that still allows another open() to read on Windows."""
    if os.name != "nt":
        return path.open("wb")
    import ctypes
    import msvcrt

    GENERIC_WRITE = 0x40000000
    FILE_SHARE_READ = 0x00000001
    FILE_SHARE_WRITE = 0x00000002
    CREATE_ALWAYS = 2
    FILE_ATTRIBUTE_NORMAL = 0x80
    kernel32 = ctypes.windll.kernel32
    kernel32.CreateFileW.restype = ctypes.c_void_p
    handle = kernel32.CreateFileW(
        str(path),
        GENERIC_WRITE,
        FILE_SHARE_READ | FILE_SHARE_WRITE,
        None,
        CREATE_ALWAYS,
        FILE_ATTRIBUTE_NORMAL,
        None,
    )
    if not handle or handle == ctypes.c_void_p(-1).value:
        raise OSError("Cannot create shared tunnel log")
    fd = msvcrt.open_osfhandle(handle, os.O_WRONLY)
    return os.fdopen(fd, "wb", buffering=0)


def _wait_local_port(timeout: float = 20.0) -> None:
    deadline = time.time() + timeout
    while time.time() < deadline:
        try:
            with socket.create_connection(("127.0.0.1", PORT), 0.6):
                return
        except OSError:
            time.sleep(0.25)


def start_public_tunnel() -> None:
    ensure_tunnel_key()
    key, _ = _tunnel_paths()
    temp_dir = Path(os.environ.get("TEMP", str(DATA_DIR)))

    def ssh_prefix() -> list[str]:
        return [
            "ssh",
            "-o", "BatchMode=yes",
            "-o", "StrictHostKeyChecking=accept-new",
            "-o", "UserKnownHostsFile=NUL",
            "-o", "GlobalKnownHostsFile=NUL",
            "-o", "TCPKeepAlive=yes",
            "-o", "ServerAliveInterval=20",
            "-o", "ServerAliveCountMax=15",
            "-o", "ExitOnForwardFailure=yes",
            "-o", "ConnectTimeout=15",
            "-o", "NumberOfPasswordPrompts=0",
            "-T",
        ]

    def build_cmd(kind: str) -> list[str]:
        prefix = ssh_prefix()
        if kind == "key":
            return prefix + [
                "-i", str(key),
                "-o", "IdentitiesOnly=yes",
                "-o", "PreferredAuthentications=publickey",
                "-R", f"80:127.0.0.1:{PORT}",
                "localhost.run",
            ]
        if kind == "nokey":
            return prefix + ["-R", f"80:127.0.0.1:{PORT}", "nokey@localhost.run"]
        return prefix + ["-p", "443", "-R", f"0:127.0.0.1:{PORT}", "free@a.pinggy.io"]

    def remember(base: str, with_key: bool) -> None:
        if not is_real_tunnel_base(base):
            print("Ignoring non-app host:", base)
            return
        PUBLIC_BASE["url"] = base.rstrip("/")
        PUBLIC_BASE["ok"] = True
        share = public_share_url()
        print("PUBLIC_URL=" + (share or ""))
        save_public_url(share)
        if with_key:
            mark_stable_tunnel(base)
            print("  ใช้คีย์ SSH แล้ว — รีสตาร์ทอาจได้ลิงก์เดิม")
        else:
            print("  ส่งลิงก์นี้ในไลน์ได้ ต้องเปิด start.bat ค้างไว้")

    def read_log(path: Path) -> str:
        try:
            return path.read_bytes().decode("utf-8", errors="replace")
        except OSError:
            return ""

    def run() -> None:
        _wait_local_port()
        kinds = ["key", "nokey", "pinggy"]
        idx = 0
        while True:
            TUNNEL_RESTART.clear()
            kind = kinds[idx % len(kinds)]
            wait_url_s = 18 if kind == "key" else 25
            print(f"Creating public HTTPS link ({kind})...")
            clear_public_link()
            log_path = temp_dir / f"pea-tunnel-{os.getpid()}-{int(time.time())}.log"
            cmd = build_cmd(kind)
            published = None
            started = time.time()
            try:
                log = _open_shared_log(log_path)
                proc = subprocess.Popen(cmd, stdout=log, stderr=subprocess.STDOUT)
                fail_count = 0
                last_health = 0.0
                while proc.poll() is None:
                    if TUNNEL_RESTART.is_set():
                        proc.terminate()
                        try:
                            proc.wait(timeout=5)
                        except Exception:
                            proc.kill()
                        break
                    try:
                        log.flush()
                    except OSError:
                        pass
                    bases = extract_tunnel_bases(read_log(log_path))
                    newest = bases[-1] if bases else None
                    if newest and newest != published:
                        remember(newest, with_key=(kind == "key"))
                        published = newest
                        last_health = time.time()
                    if not published and time.time() - started > wait_url_s:
                        print(f"ยังไม่ได้ลิงก์จาก {kind} ใน {wait_url_s} วินาที — ลองช่องทางถัดไป")
                        proc.terminate()
                        try:
                            proc.wait(timeout=5)
                        except Exception:
                            proc.kill()
                        break
                    now = time.time()
                    share = public_share_url() if PUBLIC_BASE.get("ok") else None
                    if share and now - last_health >= 25:
                        last_health = now
                        if probe_public(share):
                            fail_count = 0
                        else:
                            # Only drop the link when the remote page says the tunnel is gone.
                            try:
                                import urllib.request
                                req = urllib.request.Request(share, headers={"User-Agent": "PEA-app"})
                                with urllib.request.urlopen(req, timeout=8) as resp:
                                    body = resp.read(4000).decode("utf-8", "replace").lower()
                                if "no tunnel here" in body or "tunnel not found" in body:
                                    fail_count += 1
                                    print(f"Remote said tunnel is gone ({fail_count}) {share}")
                                    if fail_count >= 2:
                                        clear_public_link()
                                        proc.terminate()
                                        try:
                                            proc.wait(timeout=5)
                                        except Exception:
                                            proc.kill()
                                        break
                            except Exception:
                                pass
                    time.sleep(0.4)
                code = proc.wait()
                log.close()
                print("ssh tunnel exited:", code)
                log_text = read_log(log_path)
                if kind == "key" and "Permission denied" in log_text:
                    print("คีย์ SSH ยังใช้กับ localhost.run ไม่ได้ — สลับไปลิงก์ชั่วคราว")
                    idx = kinds.index("nokey")
                elif not published:
                    idx += 1
            except Exception as exc:
                print("Public tunnel error:", exc)
                idx += 1
            if TUNNEL_RESTART.is_set():
                TUNNEL_RESTART.clear()
                idx = 0
                continue
            if not PUBLIC_BASE.get("ok"):
                clear_public_link()
            print("Public link dropped, reconnecting in 3s...")
            time.sleep(3)

    threading.Thread(target=run, daemon=True, name="public-tunnel").start()


def prevent_windows_sleep() -> None:
    """Keep the PC awake while start.bat is running.

    Screen-off on many Windows laptops enters Modern Standby and kills the
    SSH tunnel. Blocking both sleep and display timeout avoids that.
    """
    if os.name != "nt":
        return
    import ctypes

    es_continuous = 0x80000000
    es_system_required = 0x00000001
    es_display_required = 0x00000002
    es_awaymode_required = 0x00000040
    flags = es_continuous | es_system_required | es_display_required | es_awaymode_required
    kernel32 = ctypes.windll.kernel32

    def pulse() -> None:
        while True:
            kernel32.SetThreadExecutionState(flags)
            time.sleep(30)

    kernel32.SetThreadExecutionState(flags)
    threading.Thread(target=pulse, daemon=True, name="stay-awake").start()
    print("  กันเครื่องหลับและกันจอดับอัตโนมัติระหว่างเปิด start.bat")
    print("  ห้ามปิดฝาโน้ตบุ๊ค / ห้ามกด Sleep — ปิดหน้าต่างนี้เมื่อเลิกใช้")


def print_banner() -> None:
    cfg = load_config()
    ip = lan_ip()
    path = f"/f/{cfg['token']}"
    print()
    print("=" * 60)
    print("  PEA Data Collection  |  ระบบเก็บข้อมูลหม้อแปลง")
    print("=" * 60)
    print(f"  เปิดบนเครื่องนี้ : http://127.0.0.1:{PORT}{path}")
    print(f"  ลิงก์ใน Wi-Fi    : http://{ip}:{PORT}{path}")
    print("  ลิงก์สาธารณะใช้ได้เฉพาะตอนเปิด start.bat ค้างไว้บนโน้ตบุ๊ค")
    print("  กันเครื่องหลับอัตโนมัติแล้ว — ห้ามปิดฝา / ห้ามกด Sleep")
    print("  ส่งเฉพาะลิงก์ที่มี lhr.life หรือ pinggy.link")
    print("=" * 60)
    print()


ensure_dirs()
load_config()
try:
    print(f"Loaded {len(sync_locations())} transformer locations")
except Exception as exc:
    print("Location load warning:", exc)

if __name__ == "__main__":
    print_banner()
    if is_hosted():
        print("Cloud host detected — skipping local tunnel")
    else:
        prevent_windows_sleep()
        start_public_tunnel()
    app.run(host=HOST, port=PORT, debug=False, threaded=True)
